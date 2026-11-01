"""TC 목록(tc/<게임>_tc.json)을 엑셀 TC 파일(tc/<게임>_TC.xlsx)로 저장한다.

사용:
  python -I tools/export_tc_xlsx.py                 # 슬라임 게임
  python -I tools/export_tc_xlsx.py --game bloodfall
  python -I tools/export_tc_xlsx.py --all
  (make_report.py 기본 실행 때도 자동으로 불린다)

- 슬라임 게임: LKM이 쓴 원본 엑셀(TC/슬라임게임_TC.xlsx)을 그대로 복사하고 오른쪽에 자동화 칸을 붙인다. 원본 파일은 고치지 않는다.
- BLOODFALL: JSON의 TC 초안으로 같은 양식의 엑셀을 새로 만든다.
- 결과·실행 시각·Jira 칸은 매번 다시 쓴다. 최근 결과 파일에 그 게임 결과가 없으면 이전 엑셀의 값을 남긴다.
- '검토' 칸은 엑셀에서 직접 바꿔도 된다(초안/확인함/버림). 다음 실행 때 JSON에 옮겨 적는다.
- openpyxl(무료 오픈소스) 필요: pip install openpyxl
"""
import json
import os
import sys
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (GAMES, LABEL, RESULTS_FILE, ROOT, TC_ID_RE, load_jira_state,  # noqa: E402
                    load_results, load_tcs, local_time, results_by_tc)

REVIEW_VALUES = ("초안", "확인함", "버림")
# 원본 TC 양식 뒤에 붙이는 자동화 칸 (이름, 너비)
AUTO_COLS = [("TC ID", 11), ("자동화", 8), ("자동화하지 않은 이유", 36), ("최근 결과", 10),
             ("실행 시각", 17), ("Jira TC", 11), ("Jira 버그", 11), ("검토", 9)]

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="FFC000")      # 원본 엑셀 머리글과 같은 색
AUTO_HEAD_FILL = PatternFill("solid", fgColor="BDD7EE")  # 자동화 칸은 다른 색으로 구분
HEAD_FONT = Font(name="Arial", size=11, bold=True)
BODY_FONT = Font(name="Arial", size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RESULT_STYLE = {
    "통과": (Font(name="Arial", size=10, bold=True, color="15803D"), None),
    "실패": (Font(name="Arial", size=10, bold=True, color="9C0006"), PatternFill("solid", fgColor="FFC7CE")),
    "시간 초과": (Font(name="Arial", size=10, bold=True, color="9C0006"), PatternFill("solid", fgColor="FFC7CE")),
}


def xlsx_path(game):
    return os.path.join(ROOT, "tc", f"{game}_TC.xlsx")


def expand_ids(tc_id):
    """'플레이-25~43' 같은 묶음 번호를 하나씩 푼다."""
    if "~" not in tc_id:
        return [tc_id]
    head, last = tc_id.split("~")
    prefix, first = head.rsplit("-", 1)
    return [f"{prefix}-{n}" for n in range(int(first), int(last) + 1)]


def tc_index(tcs):
    """TC 번호 → {kind, tc, reason}. 수동 TC의 '위와 같음'은 앞 이유로 바꾼다."""
    idx, prev = {}, ""
    for tc in tcs["automated"]:
        idx[tc["id"]] = {"kind": "auto", "tc": tc}
    for tc in tcs["manual"]:
        reason = prev if tc["reason"] == "위와 같음" else tc["reason"]
        prev = reason
        for i in expand_ids(tc["id"]):     # 묶음 안에 자동화한 번호(예: 플레이-26)가 있으면 자동 쪽을 남긴다
            idx.setdefault(i, {"kind": "manual", "tc": tc, "reason": reason})
    return idx


def find_header(ws):
    """'TC ID' 칸이 있는 머리글 줄과 {칸 이름: 열 번호}."""
    for row in ws.iter_rows(min_row=1, max_row=12):
        names = {c.value: c.column for c in row if isinstance(c.value, str)}
        if "TC ID" in names:
            return row[0].row, names
    return None, {}


def read_previous(path):
    """이전 엑셀에서 TC 번호별 검토·결과·실행 시각을 읽는다(사용자가 고친 검토 값을 살리기 위해)."""
    prev = {}
    if not os.path.exists(path):
        return prev
    wb = load_workbook(path)
    for ws in wb.worksheets:
        hdr, cols = find_header(ws)
        if not hdr:
            continue
        for r in range(hdr + 1, ws.max_row + 1):
            tc_id = ws.cell(r, cols["TC ID"]).value
            if tc_id:
                prev[tc_id] = {k: ws.cell(r, cols[k]).value for k in ("검토", "최근 결과", "실행 시각") if k in cols}
    return prev


def sync_reviews(game, tcs, prev):
    """엑셀에서 바뀐 검토 값을 JSON에 옮겨 적는다. 바뀐 TC 번호 목록을 돌려준다."""
    changed = []
    for kind in ("automated", "manual"):
        for tc in tcs[kind]:
            v = (prev.get(tc["id"]) or {}).get("검토")
            if "review" in tc and v in REVIEW_VALUES and v != tc["review"]:
                tc["review"] = v
                changed.append(tc["id"])
    if changed:
        path = os.path.join(ROOT, "tc", f"{game}_tc.json")
        with open(path, encoding="utf-8") as f:
            raw = json.load(f)
        reviews = {tc["id"]: tc.get("review") for k in ("automated", "manual") for tc in tcs[k]}
        for k in ("automated", "manual"):
            for tc in raw[k]:
                if tc["id"] in changed:
                    tc["review"] = reviews[tc["id"]]
        with open(path, "w", encoding="utf-8") as f:
            json.dump(raw, f, ensure_ascii=False, indent=2)
            f.write("\n")
    return changed


def auto_values(tc_id, info, by_tc, run_time, prev, jira):
    """자동화 칸에 들어갈 값 (AUTO_COLS 순서)."""
    p = prev.get(tc_id) or {}
    bugs = sorted({k for title, k in jira["bugs"].items() if tc_id in TC_ID_RE.findall(title)})
    if info is None:
        return [tc_id, "", "", "", "", "", "", ""]
    tc = info["tc"]
    if info["kind"] == "auto":
        if by_tc is None:                       # 이번 결과 파일에 이 게임이 없음 → 이전 값 유지
            result, when = p.get("최근 결과") or LABEL[None], p.get("실행 시각") or ""
        else:
            r = by_tc.get(tc_id)
            result, when = LABEL.get(r and r["status"], r and r["status"]), (run_time if r else "")
        reason = ""
    else:
        result, when, reason = "수동", "", info["reason"]
    jira_key = jira["tcs"].get(tc_id) or jira["tcs"].get(tc["id"], "")
    return [tc_id, "자동" if info["kind"] == "auto" else "수동", reason, result, when,
            jira_key, ", ".join(bugs), tc.get("review", "")]


def style_cell(c, head=False, auto=False, center=False):
    c.border = BORDER
    if head:
        c.font, c.alignment, c.fill = HEAD_FONT, CENTER, (AUTO_HEAD_FILL if auto else HEAD_FILL)
    else:
        c.font, c.alignment = BODY_FONT, (CENTER if center else LEFT)


def paint_result(c):
    font, fill = RESULT_STYLE.get(c.value, (None, None))
    if font:
        c.font = font
    if fill:
        c.fill = fill


def add_review_dropdown(ws, col, first, last):
    dv = DataValidation(type="list", formula1='"' + ",".join(REVIEW_VALUES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    letter = get_column_letter(col)
    dv.add(f"{letter}{first}:{letter}{last}")


def fit_height(ws, row, widths):
    """줄바꿈된 글자 수로 줄 높이를 대강 맞춘다(엑셀은 저장된 파일의 줄 높이를 자동으로 다시 계산하지 않음)."""
    lines = 1
    for col, w in widths.items():
        v = ws.cell(row, col).value
        if v:
            per_line = max(4, int(w / 1.9))          # 한글 한 글자 ≈ 너비 1.9
            lines = max(lines, sum(-(-len(part) // per_line) or 1 for part in str(v).split("\n")))
    ws.row_dimensions[row].height = max(20, 14 * lines + 6)


def write_auto_block(ws, hdr, start_col, rows, review_rows):
    """머리글 줄 hdr의 start_col부터 자동화 칸을 쓴다. rows = [(줄 번호, 값 목록)]."""
    for j, (name, width) in enumerate(AUTO_COLS):
        c = ws.cell(hdr, start_col + j, name)
        style_cell(c, head=True, auto=True)
        ws.column_dimensions[get_column_letter(start_col + j)].width = width
    for r, values in rows:
        for j, v in enumerate(values):
            c = ws.cell(r, start_col + j, v or None)
            style_cell(c, center=AUTO_COLS[j][0] not in ("자동화하지 않은 이유",))
            if AUTO_COLS[j][0] == "최근 결과":
                paint_result(c)
    if review_rows:
        add_review_dropdown(ws, start_col + len(AUTO_COLS) - 1, min(review_rows), max(review_rows))


# ---------------------------------------------------------------- 슬라임 게임: 원본 엑셀 + 자동화 칸

SLIME_SHEETS = ("공통", "플레이")   # 첫 시트 = 필수항목(공통-N), 둘째 시트 = 지정항목(플레이-N)


def build_slimegame(tcs, idx, by_tc, run_time, prev, jira):
    src = os.path.normpath(os.path.join(ROOT, tcs["source_xlsx"]))
    wb = load_workbook(src)
    used = set()
    for ws, prefix in zip(wb.worksheets, SLIME_SHEETS):
        hdr = next(r for r in range(1, 12) if ws.cell(r, 2).value == "번호")
        last = max(r for r in range(hdr + 1, ws.max_row + 1) if ws.cell(r, 2).value is not None)
        start_col = ws.max_column + 1
        rows = []
        for r in range(hdr + 1, last + 1):
            if ws.cell(r, 2).value is None:
                continue
            tc_id = f"{prefix}-{ws.cell(r, 2).value}"
            used.add(tc_id)
            rows.append((r, auto_values(tc_id, idx.get(tc_id), by_tc, run_time, prev, jira)))
        # 원본에 없는 탐색 TC는 이 시트 끝에 붙인다(플레이 시트)
        if prefix == "플레이":
            for tc in tcs["automated"]:
                if tc["id"].startswith("탐색"):
                    last += 1
                    for col in range(2, start_col):
                        ws.cell(last, col)._style = copy(ws.cell(last - 1, col)._style)
                    for col, v in zip(range(2, 11), (tc["id"], tc.get("priority"), tc.get("cat1"), tc.get("cat2"),
                                                     tc.get("cat3"), tc.get("pre"), tc.get("steps"),
                                                     tc.get("expected"), "자동 테스트 중 발견 — AI 추가 초안, 원본 TC에 없던 항목")):
                        ws.cell(last, col, v)
                    rows.append((last, auto_values(tc["id"], idx[tc["id"]], by_tc, run_time, prev, jira)))
                    fit_height(ws, last, {c: ws.column_dimensions[get_column_letter(c)].width or 13 for c in range(2, 11)})
                    used.add(tc["id"])
        review_rows = [r for r, v in rows if v[-1]]
        write_auto_block(ws, hdr, start_col, rows, review_rows)
        note = ws.cell(hdr - 1, start_col, "← 원본 TC (LKM)   |   자동화 칸 (export_tc_xlsx.py가 매번 다시 씀, 검토 칸만 직접 수정)")
        note.font = Font(name="Arial", size=9, italic=True, color="6B7280")
        ws.freeze_panes = ws.cell(hdr + 1, 3)
    missing = [i for i in idx if i not in used]
    return wb, missing


# ---------------------------------------------------------------- BLOODFALL: JSON 초안으로 새 엑셀

BF_COLS = [("번호", 6), ("우선순위", 9), ("분류 1", 10), ("분류 2", 14), ("확인 항목", 34), ("근거", 10),
           ("사전 조건", 26), ("테스트 내용", 38), ("테스트 기대 결과", 42), ("비고", 18)]


def build_bloodfall(tcs, idx, by_tc, run_time, prev, jira):
    wb = Workbook()
    ws = wb.active
    ws.title = "TC"
    ws.column_dimensions["A"].width = 2
    info = [("프로젝트 명칭", tcs["game"]), ("플랫폼 종류", "웹 브라우저 (데스크톱, 키보드·마우스)"),
            ("버전 정보", f"{tcs['target']} (기획서 v0.9 기준)"), ("작성", "Claude Code 초안 / 검토 LKM (검토 칸 참고)")]
    for i, (k, v) in enumerate(info, start=2):
        a, b = ws.cell(i, 9, k), ws.cell(i, 10, v)
        style_cell(a, head=True)
        style_cell(b)
        b.font = Font(name="Arial", size=11, bold=True)
    hdr = 2 + len(info)
    for j, (name, width) in enumerate(BF_COLS, start=2):
        style_cell(ws.cell(hdr, j, name), head=True)
        ws.column_dimensions[get_column_letter(j)].width = width
    start_col = 2 + len(BF_COLS)
    rows, n = [], 0
    for kind in ("automated", "manual"):
        for tc in tcs[kind]:
            n += 1
            r = hdr + n
            note = "수동 TC" if kind == "manual" else ""
            for j, v in enumerate((n, tc.get("priority"), tc.get("cat1"), tc.get("cat2"), tc.get("title"),
                                   tc.get("basis"), tc.get("pre"), tc.get("steps"), tc.get("expected"), note), start=2):
                style_cell(ws.cell(r, j, v), center=j in (2, 3, 4, 5, 7))
            rows.append((r, auto_values(tc["id"], idx[tc["id"]], by_tc, run_time, prev, jira)))
    write_auto_block(ws, hdr, start_col, rows, [r for r, v in rows if v[-1]])
    widths = {j: w for j, (_, w) in enumerate(BF_COLS + AUTO_COLS, start=2)}
    for r, _ in rows:
        fit_height(ws, r, widths)
    ws.row_dimensions[hdr].height = 30
    ws.freeze_panes = ws.cell(hdr + 1, 4)
    ws.auto_filter.ref = f"B{hdr}:{get_column_letter(start_col + len(AUTO_COLS) - 1)}{hdr + n}"
    return wb, []


BUILDERS = {"slimegame": build_slimegame, "bloodfall": build_bloodfall}


def export(game):
    path = xlsx_path(game)
    tcs = load_tcs(game)
    prev = read_previous(path)
    changed = sync_reviews(game, tcs, prev)
    by_tc, run_time = None, ""
    if os.path.exists(RESULTS_FILE):
        stats, rows = load_results(RESULTS_FILE, game)
        if rows:
            by_tc, run_time = results_by_tc(rows), local_time(stats.get("startTime"))
    wb, missing = BUILDERS[game](tcs, tc_index(tcs), by_tc, run_time, prev, load_jira_state())
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"저장 실패: {path} 가 엑셀에서 열려 있습니다. 닫고 다시 실행하세요.")
    print(f"TC 엑셀 저장: {path}"
          + (f" (결과: {run_time} 실행분)" if by_tc else " (최근 결과 파일에 이 게임 결과 없음 → 이전 결과 유지)"))
    if changed:
        print(f"  엑셀에서 바뀐 검토 값을 JSON에 반영: {', '.join(changed)}")
    if missing:
        print(f"  경고: 엑셀에서 찾지 못한 TC 번호 {missing}")
    import check_tc                                  # TC 작성 규칙(044) 점검 한 줄 요약
    print("  " + check_tc.summary(game))
    return path


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    games = GAMES if "--all" in sys.argv else [GAMES[0] if "--game" not in sys.argv
                                                else sys.argv[sys.argv.index("--game") + 1]]
    for g in games:
        if g not in GAMES:
            sys.exit(f"--game 은 {', '.join(GAMES)} 중 하나")
        export(g)


if __name__ == "__main__":
    main()
