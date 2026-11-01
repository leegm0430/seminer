"""TC를 'TC 작성 규칙'(AX부트캠프study/Resources/2026-11-02 044. TC 작성 규칙.md)에 맞는지 점검한다.

기계로 볼 수 있는 규칙만 본다. 문장이 자연스러운지, 빠진 케이스가 없는지는 사람이 본다.

사용:
  python -I tools/check_tc.py --game bloodfall                 # qa-automation TC 목록(JSON, 없으면 생성된 엑셀)
  python -I tools/check_tc.py --xlsx "../TC/운빨존많겜TC작성자_이경민.xlsx"   # 직접 쓴 엑셀 TC
  python -I tools/check_tc.py --all                            # qa-automation 게임 전부
"""
import os
import re
import sys
from collections import Counter, defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import GAMES, ROOT, load_tcs  # noqa: E402

RULES = {
    "S4": "사전 조건·테스트 내용·기대 결과 중 빈 칸",
    "W4-끝맺음": "기대 결과가 '확인/되는지'로 끝남 → 결과(~ 됨)로 쓴다",
    "W4-정상": "기대 결과가 '정상/제대로/문제없이'뿐 → 무엇이 보이는지 쓴다",
    "Q1": "기대 결과가 테스트 내용과 똑같음",
    "D6": "확률·무작위 TC에 시행 횟수(N회)가 없음",
    "Q2": "번호가 비었거나 겹침, 또는 번호만 있고 내용이 없음",
    "P1": "High(높음) 비율이 80% 넘음 → 우선순위 정의대로 다시 보기",
    "Q4": "우선순위 표기가 섞임 (High/Medium/Low 와 높음/보통/낮음)",
}
END_CHECK = re.compile(r"(확인|되는지|인지|는지)\s*[.)\]]?\s*$")
VAGUE = re.compile(r"정상|제대로|문제\s*없|이상\s*없")
# "정상" 같은 말과 흔한 서술어를 빼고 남는 내용이 거의 없으면 막연한 결과로 본다
FILLER = re.compile(r"정상적으로|정상|제대로|문제없이|이상없이|작동|동작|진행|실행|됨|된다|되어야한다|이|가|을|를|\.")
RANDOM = re.compile(r"확률(에\s*따른|대로|로)|랜덤(하게|성)|무작위로|분포")   # 결과가 무작위인지·확률대로인지 판정하는 TC
TRIALS = re.compile(r"\d[\d,]*\s*(회|번)")
HIGH = {"high", "높음"}
EN, KO = {"high", "medium", "low"}, {"높음", "보통", "낮음"}


def norm(s):
    return re.sub(r"\s+", "", str(s or ""))


def check_rows(rows):
    """rows: [{sheet, id, priority, pre, steps, expected}] → {규칙: [(sheet, id), ...]}"""
    found = defaultdict(list)
    by_sheet = defaultdict(list)
    for r in rows:
        by_sheet[r["sheet"]].append(r)
    for sheet, items in by_sheet.items():
        dup = {i for i, n in Counter(r["id"] for r in items).items() if i and n > 1}
        last_pre = ""
        for r in items:
            where = (sheet, r["id"] or "(번호 없음)")
            steps, exp = norm(r["steps"]), norm(r["expected"])
            if not steps and not exp:
                found["Q2"].append(where)          # 번호만 있는 빈 줄
                continue
            if not r["id"] or r["id"] in dup:
                found["Q2"].append(where)
            pre = r["pre"] or last_pre               # 엑셀은 위 칸과 같은 사전 조건을 비워 두기도 한다(병합)
            last_pre = pre
            if not norm(pre) or not steps or not exp:
                found["S4"].append(where)
            if exp and END_CHECK.search(str(r["expected"]).strip()):
                found["W4-끝맺음"].append(where)
            if exp and VAGUE.search(str(r["expected"])) and len(FILLER.sub("", exp)) < 6:
                found["W4-정상"].append(where)
            if exp and exp == steps:
                found["Q1"].append(where)
            text = f"{r['steps']} {r['expected']}"
            if RANDOM.search(text) and not TRIALS.search(text):
                found["D6"].append(where)
        prios = [str(r["priority"]).strip().lower() for r in items if r["priority"]]
        highs = sum(p in HIGH for p in prios)
        if len(prios) >= 10 and highs / len(prios) > 0.8:
            found["P1"].append((sheet, f"{highs}/{len(prios)}건이 High"))
        if set(prios) & EN and set(prios) & KO:
            found["Q4"].append((sheet, "표기 섞임"))
    return found


# ---------------------------------------------------------------- 읽기

HEADER_KEYS = {
    "id": ("tc_id", "tc id"),
    "no": ("no.", "번호", "tc_no"),
    "priority": ("우선", "priority"),
    "pre": ("사전",),
    "steps": ("테스트 내용", "test details", "절차"),
    "expected": ("기대", "expected"),
}


def _match(text, keys):
    t = str(text or "").lower().replace("\n", " ")
    return any(k in t for k in keys)


def rows_from_xlsx(path):
    """TC 엑셀을 읽는다. 머리글 이름으로 칸을 찾으므로 양식이 조금 달라도 된다(머리글 두 줄도 됨)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        hdr = next((r for r in range(1, 16) if any(_match(c.value, HEADER_KEYS["expected"]) for c in ws[r])), None)
        if not hdr:
            continue                                   # 표지 시트 등
        cols, last_hdr = {}, hdr
        for r in (hdr, hdr + 1):                       # 머리글이 두 줄인 양식(Weird Space)
            hit = False
            for c in ws[r]:
                for k, keys in HEADER_KEYS.items():
                    if k not in cols and _match(c.value, keys):
                        cols[k], hit = c.column, True
            if hit:
                last_hdr = r
        for r in range(last_hdr + 1, ws.max_row + 1):
            def get(k):
                v = ws.cell(r, cols[k]).value if k in cols else None
                return None if isinstance(v, str) and not v.strip() else v
            num = get("id") or get("no")
            if num is None and not any(get(k) for k in ("pre", "steps", "expected")):
                continue
            if isinstance(num, str) and "공란" in num:
                continue
            rows.append({"sheet": ws.title.strip(), "id": str(num).strip() if num is not None else "",
                         "priority": get("priority"), "pre": get("pre") or "",
                         "steps": get("steps") or "", "expected": get("expected") or ""})
    return rows


def rows_from_json(game):
    tcs = load_tcs(game)
    items = tcs["automated"] + tcs["manual"]
    if not all(t.get("expected") for t in items):
        return None                                    # 내용 칸이 JSON에 없음(슬라임 → 엑셀에서 읽기)
    return [{"sheet": game, "id": t["id"], "priority": t.get("priority"), "pre": t.get("pre", ""),
             "steps": t.get("steps", ""), "expected": t.get("expected", "")} for t in items]


def rows_for_game(game):
    rows = rows_from_json(game)
    if rows is None:
        path = os.path.join(ROOT, "tc", f"{game}_TC.xlsx")
        rows = rows_from_xlsx(path) if os.path.exists(path) else []
    return rows


# ---------------------------------------------------------------- 출력

def report(name, rows, limit=12):
    found = check_rows(rows)
    total = sum(len(v) for v in found.values())
    print(f"\n[{name}] TC {len(rows)}줄 점검 → 규칙에 걸린 곳 {total}건")
    for rule, label in RULES.items():
        hits = found.get(rule)
        if not hits:
            continue
        by_sheet = defaultdict(list)
        for sheet, i in hits:
            by_sheet[sheet].append(str(i))
        lines = []
        for sheet, ids in by_sheet.items():
            shown = ", ".join(ids[:limit]) + (f" 외 {len(ids) - limit}건" if len(ids) > limit else "")
            lines.append(f"{sheet}: {shown}")
        print(f"  {rule} {label} ({len(hits)}건)\n     " + "\n     ".join(lines))
    if total == 0:
        print("  걸린 곳 없음. (문장 품질·빠진 케이스는 사람이 따로 본다)")
    return found


def summary(game):
    """export_tc_xlsx.py가 쓰는 한 줄 요약."""
    found = check_rows(rows_for_game(game))
    total = sum(len(v) for v in found.values())
    if not total:
        return "TC 규칙 점검: 걸린 곳 없음"
    detail = ", ".join(f"{k} {len(v)}" for k, v in found.items())
    return f"TC 규칙 점검: {total}건 ({detail}) → python -I tools/check_tc.py --game {game}"


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    args = sys.argv[1:]
    if "--xlsx" in args:
        path = args[args.index("--xlsx") + 1]
        report(os.path.basename(path), rows_from_xlsx(path))
    else:
        games = GAMES if "--all" in args else [args[args.index("--game") + 1] if "--game" in args else GAMES[0]]
        for g in games:
            report(g, rows_for_game(g))
    print("\n규칙 설명: AX부트캠프study/Resources/2026-11-02 044. TC 작성 규칙.md")


if __name__ == "__main__":
    main()
